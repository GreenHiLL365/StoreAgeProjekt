from flask import Flask, render_template, request, redirect, url_for, flash, Response
import mysql.connector
from openpyxl import load_workbook
import xlrd
from datetime import datetime
import re
import io
import csv

app = Flask(__name__)
app.secret_key = "hemmelig"

db_config = {
    'host': 'localhost',
    'user': 'root',
    'password': 'Wgrgryjr3@',
    'database': 'glaslager'
}

MONTH_MAP = {
    "januar": 1, "februar": 2, "marts": 3, "april": 4,
    "maj": 5, "juni": 6, "juli": 7, "august": 8,
    "september": 9, "oktober": 10, "november": 11, "december": 12
}


materialer = [
    "kul_sække", "kul_bigbag",
    "jernoxyd", "koboltoxyd", "pyrit",
    "salpeter_bigbag", "zinkselenit", "chromit"
]
felter = ["pr_palle", "hele_paller", "loese", "pr_enhed", "total_enheder", "total_kg"]

celle_mapping = {
    "optalt_af": "B22",
    "registeret_maaned": "B25",  # maaned fra Excel
    "kul_sække_pr_palle": "B13", "kul_sække_hele_paller": "C13", "kul_sække_loese": "D13", "kul_sække_pr_enhed": "E13", "kul_sække_total_enheder": "F13", "kul_sække_total_kg": "G13",
    "kul_bigbag_pr_palle": "B14", "kul_bigbag_hele_paller": "C14", "kul_bigbag_loese": "D14", "kul_bigbag_pr_enhed": "E14", "kul_bigbag_total_enheder": "F14", "kul_bigbag_total_kg": "G14",
    "jernoxyd_pr_palle": "B15", "jernoxyd_hele_paller": "C15", "jernoxyd_loese": "D15", "jernoxyd_pr_enhed": "E15", "jernoxyd_total_enheder": "F15", "jernoxyd_total_kg": "G15",
    "koboltoxyd_pr_palle": "B16", "koboltoxyd_hele_paller": "C16", "koboltoxyd_loese": "D16", "koboltoxyd_pr_enhed": "E16", "koboltoxyd_total_enheder": "F16", "koboltoxyd_total_kg": "G16",
    "pyrit_pr_palle": "B17", "pyrit_hele_paller": "C17", "pyrit_loese": "D17", "pyrit_pr_enhed": "E17", "pyrit_total_enheder": "F17", "pyrit_total_kg": "G17",
    "salpeter_bigbag_pr_palle": "B18", "salpeter_bigbag_hele_paller": "C18", "salpeter_bigbag_loese": "D18", "salpeter_bigbag_pr_enhed": "E18", "salpeter_bigbag_total_enheder": "F18", "salpeter_bigbag_total_kg": "G18",
    "zinkselenit_pr_palle": "B19", "zinkselenit_hele_paller": "C19", "zinkselenit_loese": "D19", "zinkselenit_pr_enhed": "E19", "zinkselenit_total_enheder": "F19", "zinkselenit_total_kg": "G19",
    "chromit_pr_palle": "B20", "chromit_hele_paller": "C20", "chromit_loese": "D20", "chromit_pr_enhed": "E20", "chromit_total_enheder": "F20", "chromit_total_kg": "G20",
}

INT_FIELDS = {"hele_paller", "loese", "total_enheder"}
FLOAT_FIELDS = {"pr_enhed", "total_kg", "pr_palle"}

def parse_number(value):
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return value
    s = str(value).strip().replace('\u00A0','').replace(' ','')
    if s == "":
        return None
    if ',' in s and '.' in s:
        s = s.replace('.', '').replace(',', '.')
    else:
        if ',' in s and '.' not in s:
            s = s.replace(',', '.')
        if s.count('.') > 1:
            s = s.replace('.', '')
    s_clean = re.sub(r'[^0-9\.\-]', '', s)
    if s_clean in ("", "-", "."):
        return None
    try:
        return float(s_clean) if '.' in s_clean else int(s_clean)
    except:
        try:
            return float(s_clean)
        except:
            return None


@app.route("/", methods=["GET","POST"])
def index():
    conn = mysql.connector.connect(**db_config)
    cursor = conn.cursor(dictionary=True)
    lager = {}
    dato_test = None

    # Upload Excel
    if request.method == "POST" and "upload_file" in request.form:
        file = request.files.get("excel_file")
        if file:
            filename = file.filename.lower()
            try:
                if filename.endswith(".xlsx"):
                    wb = load_workbook(file)
                    sheet = wb.active
                    for key, celle in celle_mapping.items():
                        try:
                            val = sheet[celle].value
                            if any(key.endswith(f) for f in felter):
                                val = parse_number(val)
                            lager[key] = val
                        except:
                            lager[key] = None
                    # Hent rå dato fra Excel
                    raw_dato = sheet["B23"].value
                    lager["raa_dato"] = raw_dato

                elif filename.endswith(".xls"):
                    book = xlrd.open_workbook(file_contents=file.read())
                    sheet = book.sheet_by_index(0)
                    for key, celle in celle_mapping.items():
                        col_letter = ''.join(filter(str.isalpha, celle))
                        row_number = int(''.join(filter(str.isdigit, celle)))
                        col_index = ord(col_letter.upper()) - 65
                        row_index = row_number - 1
                        try:
                            val = sheet.cell_value(row_index, col_index)
                            if any(key.endswith(f) for f in felter):
                                val = parse_number(val)
                            lager[key] = val
                        except:
                            lager[key] = None
                    # Hent rå dato fra Excel
                    raw_dato = sheet.cell_value(22, 1)
                    lager["raa_dato"] = raw_dato

                # Konverter Excel-dato til <input type="date">
                optalt_dato = None
                if isinstance(raw_dato, datetime):
                    optalt_dato = raw_dato.date()
                else:
                    try:
                        optalt_dato = datetime.strptime(str(raw_dato), "%d.%m.%Y").date()
                    except:
                        optalt_dato = None
                if optalt_dato:
                    lager["optalt_dato"] = optalt_dato.strftime("%Y-%m-%d")
                else:
                    lager["optalt_dato"] = ""

                flash("Excel-data indlæst. Kontroller felterne før gem.")
            except Exception as e:
                flash(f"Fejl ved upload: {e}")
        else:
            flash("Ingen fil valgt")

    # Gem formular
    if request.method == "POST" and "save_form" in request.form:
        data = {}
        # Optællingsdato
        optalt_dato = request.form.get("optalt_dato") or lager.get("optalt_dato")
        data["optalt_dato"] = optalt_dato

        # Måned
        raw_month = request.form.get("registeret_maaned") or lager.get("registeret_maaned")
        if raw_month:
            raw_month_str = str(raw_month).strip().lower()
            if raw_month_str.isdigit():
                data["registeret_maaned"] = int(raw_month_str)
            else:
                data["registeret_maaned"] = MONTH_MAP.get(raw_month_str, None)
        else:
            data["registeret_maaned"] = None

        # Optalt af og år
        data["optalt_af"] = request.form.get("optalt_af") or lager.get("optalt_af")
        data["registeret_aar"] = request.form.get("registeret_aar") or lager.get("registeret_aar")

        # Materialer
        for m in materialer:
            for f in ["pr_palle","hele_paller","loese","pr_enhed","total_enheder","total_kg"]:
                field_name = f"{m}_{f}"
                raw = request.form.get(field_name)
                if raw in (None,""):
                    raw = lager.get(field_name)
                if raw in (None,""):
                    data[field_name] = None
                    continue
                if f in {"hele_paller","loese","total_enheder"}:
                    parsed = parse_number(raw)
                    data[field_name] = int(parsed) if parsed is not None else None
                elif f in {"pr_enhed","total_kg","pr_palle"}:
                    parsed = parse_number(raw)
                    data[field_name] = float(parsed) if parsed is not None else None
                else:
                    data[field_name] = raw

        # Beregn totaler hvis nødvendigt
        for m in materialer:
            try:
                pr_palle = float(data.get(f"{m}_pr_palle") or 0)
                hele_paller = int(data.get(f"{m}_hele_paller") or 0)
                loese = int(data.get(f"{m}_loese") or 0)
                pr_enhed = float(data.get(f"{m}_pr_enhed") or 0)
                if not data.get(f"{m}_total_enheder"):
                    total_enheder = hele_paller*pr_palle + loese
                    data[f"{m}_total_enheder"] = int(total_enheder) if float(total_enheder).is_integer() else float(total_enheder)
                if not data.get(f"{m}_total_kg"):
                    total_kg = float(data[f"{m}_total_enheder"])*pr_enhed
                    data[f"{m}_total_kg"] = float(total_kg)
            except:
                pass

        # Gem i DB
        cols = ", ".join(data.keys())
        placeholders = ", ".join(["%s"]*len(data))
        sql = f"INSERT INTO Lager ({cols}) VALUES ({placeholders})"
        cursor.execute(sql,list(data.values()))
        conn.commit()
        cursor.close()
        conn.close()
        flash("Data gemt i databasen")
        return redirect(url_for("index"))

    if not lager:
        cursor.execute("SELECT * FROM Lager ORDER BY id DESC LIMIT 1")
        lager = cursor.fetchone() or {}

    return render_template("index.html", lager=lager, materialer=materialer, dato_test=dato_test)


if __name__ == "__main__":
    app.run(debug=True)
