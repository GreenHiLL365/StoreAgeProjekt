from flask import Flask, render_template, request, redirect, url_for, flash, Response
import mysql.connector
from openpyxl import load_workbook
import xlrd
from datetime import datetime
import calendar
import re
import io
import csv

app = Flask(__name__)
app.secret_key = "hemmelig"

# MySQL konfiguration
db_config = {
    'host': 'localhost',
    'user': 'root',
    'password': 'Wgrgryjr3@',
    'database': 'glaslager'
}

# Materialer og felter
materialer = [
    "kul_sække", "kul_bigbag",
    "jernoxyd", "koboltoxyd", "pyrit",
    "salpeter_bigbag", "zinkselenit", "chromit"
]
felter = ["pr_palle", "hele_paller", "loese", "pr_enhed", "total_enheder", "total_kg"]

# CELLE MAPPING for Excel (undtagen dato)
celle_mapping = {
    "optalt_af": "B22",
    "kul_sække_pr_palle": "B13", "kul_sække_hele_paller": "C13", "kul_sække_loese": "D13", "kul_sække_pr_enhed": "E13", "kul_sække_total_enheder": "F13", "kul_sække_total_kg": "G13",
    "kul_bigbag_pr_palle": "B14", "kul_bigbag_hele_paller": "C14", "kul_bigbag_loese": "D14", "kul_bigbag_pr_enhed": "E14", "kul_bigbag_total_enheder": "F14", "kul_bigbag_total_kg": "G14",
    "jernoxyd_pr_palle": "B15", "jernoxyd_hele_paller": "C15", "jernoxyd_loese": "D15", "jernoxyd_pr_enhed": "E15", "jernoxyd_total_enheder": "F15", "jernoxyd_total_kg": "G15",
    "koboltoxyd_pr_palle": "B16", "koboltoxyd_hele_paller": "C16", "koboltoxyd_loese": "D16", "koboltoxyd_pr_enhed": "E16", "koboltoxyd_total_enheder": "F16", "koboltoxyd_total_kg": "G16",
    "pyrit_pr_palle": "B17", "pyrit_hele_paller": "C17", "pyrit_loese": "D17", "pyrit_pr_enhed": "E17", "pyrit_total_enheder": "F17", "pyrit_total_kg": "G17",
    "salpeter_bigbag_pr_palle": "B18", "salpeter_bigbag_hele_paller": "C18", "salpeter_bigbag_loese": "D18", "salpeter_bigbag_pr_enhed": "E18", "salpeter_bigbag_total_enheder": "F18", "salpeter_bigbag_total_kg": "G18",
    "zinkselenit_pr_palle": "B19", "zinkselenit_hele_paller": "C19", "zinkselenit_loese": "D19", "zinkselenit_pr_enhed": "E19", "zinkselenit_total_enheder": "F19", "zinkselenit_total_kg": "G19",
    "chromit_pr_palle": "B20", "chromit_hele_paller": "C20", "chromit_loese": "D20", "chromit_pr_enhed": "E20", "chromit_total_enheder": "F20", "chromit_total_kg": "G20",
}

INT_FIELDS = {"hele_paller", "loese", "total_enheder"}  # tæller/antal
FLOAT_FIELDS = {"pr_enhed", "total_kg", "pr_palle"}     # vægt/pr enhed mv. (pr_palle kan være float hvis nødvendigt)


def parse_number(value):
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return value
    s = str(value).strip()
    if s == "":
        return None
    s = s.replace('\u00A0', '').replace(' ', '')
    if ',' in s and '.' in s:
        s = s.replace('.', '').replace(',', '.')
    else:
        if ',' in s and '.' not in s:
            s = s.replace(',', '.')
        if s.count('.') > 1:
            s = s.replace('.', '')
    s_clean = re.sub(r'[^0-9\.\-]', '', s)
    if s_clean in ("", "-", "."):
        return None
    try:
        if '.' in s_clean:
            return float(s_clean)
        else:
            return int(s_clean)
    except:
        try:
            return float(s_clean)
        except:
            return None


@app.route("/", methods=["GET", "POST"])
def index():
    conn = mysql.connector.connect(**db_config)
    cursor = conn.cursor(dictionary=True)
    lager = {}
    dato_test = None  # rå data fra Excel

    # Upload Excel
    if request.method == "POST" and "upload_file" in request.form:
        file = request.files.get("excel_file")
        if file:
            filename = file.filename.lower()
            try:
                if filename.endswith(".xlsx"):
                    wb = load_workbook(file)
                    sheet = wb.active
                    for key, celle in celle_mapping.items():
                        try:
                            val = sheet[celle].value
                            if any(key.endswith(f) for f in ["pr_palle", "pr_enhed", "total_enheder", "total_kg", "hele_paller", "loese"]):
                                val = parse_number(val)
                            lager[key] = val
                        except:
                            lager[key] = None
                    raw_dato = sheet["B23"].value

                elif filename.endswith(".xls"):
                    book = xlrd.open_workbook(file_contents=file.read())
                    sheet = book.sheet_by_index(0)
                    for key, celle in celle_mapping.items():
                        col_letter = ''.join(filter(str.isalpha, celle))
                        row_number = int(''.join(filter(str.isdigit, celle)))
                        col_index = ord(col_letter.upper()) - 65
                        row_index = row_number - 1
                        try:
                            val = sheet.cell_value(row_index, col_index)
                            if any(key.endswith(f) for f in ["pr_palle", "pr_enhed", "total_enheder", "total_kg", "hele_paller", "loese"]):
                                val = parse_number(val)
                            lager[key] = val
                        except:
                            lager[key] = None
                    raw_dato = sheet.cell_value(22, 1)

                else:
                    flash("Vælg venligst en gyldig Excel-fil (.xlsx eller .xls)")
                    return redirect(url_for("index"))

                # Dato-opdeling
                if raw_dato:
                    try:
                        if isinstance(raw_dato, datetime):
                            dag = raw_dato.day
                            maaned = raw_dato.month
                            aar = raw_dato.year
                        else:
                            parts = str(raw_dato).replace("-", ".").split(".")
                            if len(parts) >= 3:
                                dag = parts[0]
                                maaned_part = parts[1]
                                aar = parts[2]
                                try:
                                    maaned = int(maaned_part)
                                except ValueError:
                                    maaned = list(calendar.month_name).index(maaned_part.capitalize())
                            else:
                                raise ValueError("Ugyldigt datoformat")
                        # gem som dd.mm.yyyy og måned som 2 cifre
                        lager["optalt_dato"] = f"{int(dag):02}.{int(maaned):02}.{int(aar)}"
                        lager["optalt_maaned"] = f"{int(maaned):02}"
                        lager["optalt_aar"] = str(aar)
                        dato_test = str(raw_dato)  # rå data til visning
                    except Exception as e:
                        lager["optalt_dato"] = None
                        dato_test = f"Ugyldig dato: {raw_dato} indtast manuelt ({e})"

                flash("Excel-data indlæst. Kontroller felterne før gem.")
            except Exception as e:
                flash(f"Fejl ved upload: {e}")
        else:
            flash("Ingen fil valgt")

    # Gem formular-data
    if request.method == "POST" and "save_form" in request.form:
        data = {}
        data["optalt_dato"] = request.form.get("optalt_dato") or None
        data["optalt_maaned"] = request.form.get("optalt_maaned") or None
        data["optalt_aar"] = request.form.get("optalt_aar") or None
        data["optalt_af"] = request.form.get("optalt_af") or None

        for m in materialer:
            for f in felter:
                field_name = f"{m}_{f}"
                raw = request.form.get(field_name)
                if raw is None or raw == "":
                    data[field_name] = None
                    continue
                if f in INT_FIELDS:
                    parsed = parse_number(raw)
                    data[field_name] = int(parsed) if parsed is not None else None
                elif f in FLOAT_FIELDS:
                    parsed = parse_number(raw)
                    data[field_name] = float(parsed) if parsed is not None else None
                else:
                    parsed = parse_number(raw)
                    if isinstance(parsed, (int, float)):
                        data[field_name] = parsed
                    else:
                        data[field_name] = raw

            try:
                pr_palle = float(data.get(f"{m}_pr_palle") or 0)
                hele_paller = int(data.get(f"{m}_hele_paller") or 0)
                loese = int(data.get(f"{m}_loese") or 0)
                pr_enhed = float(data.get(f"{m}_pr_enhed") or 0)

                if not data.get(f"{m}_total_enheder"):
                    total_enheder = hele_paller * pr_palle + loese
                    data[f"{m}_total_enheder"] = int(total_enheder) if float(total_enheder).is_integer() else float(total_enheder)
                if not data.get(f"{m}_total_kg"):
                    total_kg = float(data[f"{m}_total_enheder"]) * pr_enhed
                    data[f"{m}_total_kg"] = float(total_kg)
            except Exception as e:
                pass

        cols = ", ".join(data.keys())
        placeholders = ", ".join(["%s"] * len(data))
        sql = f"INSERT INTO Lager ({cols}) VALUES ({placeholders})"
        try:
            cursor = conn.cursor()
        except:
            conn = mysql.connector.connect(**db_config)
            cursor = conn.cursor()
        cursor.execute(sql, list(data.values()))
        conn.commit()
        cursor.close()
        conn.close()
        flash("Data gemt i databasen")
        return redirect(url_for("index"))

    if not lager:
        cursor.execute("SELECT * FROM Lager ORDER BY id DESC LIMIT 1")
        lager = cursor.fetchone() or {}

    for key, value in list(lager.items()):
        if isinstance(value, (int, float)):
            if isinstance(value, int):
                lager[key] = f"{value:,}".replace(",", ".")
            else:
                lager[key] = f"{value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")

    return render_template("index.html", lager=lager, materialer=materialer, dato_test=dato_test)


@app.route("/upload_csv_example", methods=["GET"])
def download_example():
    output = io.StringIO()
    writer = csv.writer(output, delimiter=';')
    writer.writerow(["optalt_aar", "optalt_maaned", "optalt_dato", "optalt_af", "kul_sække_pr_palle", "kul_sække_total_kg"])
    writer.writerow([2025, 9, "23.09.2025", "Mabba", 20, 1234.5])
    output.seek(0)
    return Response(output.getvalue(), mimetype="text/csv", headers={"Content-Disposition":"attachment; filename=example.csv"})


if __name__ == "__main__":
    app.run(debug=True)
